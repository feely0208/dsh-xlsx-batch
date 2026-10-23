#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
dsh-xlsx-batch 的真实处理逻辑（openpyxl）

由 bin/cli.mjs 通过 spawnSync 调用：
    python3 lib/xlsx_tool.py <action> '<json payload>'
约定：
  - stdout 只输出一行 JSON（最后一行），人类可读文案由 Node 侧负责
  - stderr 输出异常栈（调试用）
  - 退出码：0 成功 / 1 运行失败 / 2 参数或输入问题 / 3 缺依赖
  - 只读输入文件，所有产物写进 --out 目录（默认 <输入目录>/_输出）
  - 读取用 data_only=True：公式取 Excel 缓存的计算结果；输出为值（不是公式）
"""

import csv
import json
import os
import re
import sys
import tempfile
import traceback
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
except ImportError as exc:  # 缺依赖必须明确提示安装命令，不能静默
    print(json.dumps({
        "ok": False,
        "code": 3,
        "message": "缺少 Python 依赖：%s" % getattr(exc, "name", exc),
        "hint": "请先安装：\n  python3 -m pip install openpyxl\n"
                "如果 python3 也没有：\n  macOS: brew install python3\n  Ubuntu: sudo apt install python3",
    }, ensure_ascii=False))
    sys.exit(3)

XLSX_EXTS = {".xlsx", ".xlsm"}
SKIP_DIRS = {"_输出", "node_modules", ".git", "__pycache__", ".venv", "venv"}
SOURCE_FILE = "来源文件"
SOURCE_SHEET = "来源工作表"
NUM_RE = re.compile(r"^-?\d{1,3}(?:,\d{3})*(?:\.\d+)?$|^-?\d+(?:\.\d+)?$")


# ── 错误 ──────────────────────────────────────────────────────────

class ToolError(Exception):
    def __init__(self, message, hint="", code=2):
        super().__init__(message)
        self.message = message
        self.hint = hint
        self.code = code


def emit(obj):
    print(json.dumps(obj, ensure_ascii=False))


# ── 路径与文件收集 ────────────────────────────────────────────────

def natural_key(text):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", text)]


def resolve_out(payload, input_path):
    out = payload.get("out")
    if out:
        return Path(out).expanduser().resolve()
    base = Path(input_path).expanduser().resolve()
    if not base.is_dir():
        base = base.parent
    return base / "_输出"


def collect_xlsx(input_path, out_dir=None):
    if not input_path:
        raise ToolError("缺少 --input", "例：merge --input ./报表目录 --out ./汇总")
    root = Path(input_path).expanduser().resolve()
    if not root.exists():
        raise ToolError(
            "输入不存在：%s" % input_path,
            "请检查路径是否正确；传目录会递归处理其中的 .xlsx 文件。",
        )
    skip_out = set()
    if out_dir:
        try:
            skip_out = {str(out_dir.resolve())}
        except OSError:
            skip_out = set()
    if root.is_file():
        if root.suffix.lower() not in XLSX_EXTS:
            raise ToolError(
                "输入不是 Excel 文件：%s" % root.name,
                "本插件只处理 .xlsx / .xlsm；老格式 .xls 请先另存为 .xlsx。",
            )
        return [root]
    files = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = sorted(
            [d for d in dirnames if d not in SKIP_DIRS and not d.startswith("~$")],
            key=natural_key,
        )
        if str(Path(dirpath).resolve()) in skip_out:
            dirnames[:] = []
            continue
        for name in filenames:
            if name.startswith("~$"):
                continue  # Excel 打开时的临时文件
            p = Path(dirpath) / name
            if p.suffix.lower() in XLSX_EXTS:
                files.append(p)
    files.sort(key=lambda p: natural_key(str(p.relative_to(root))))
    if not files:
        raise ToolError(
            "目录里没有 .xlsx / .xlsm 文件：%s" % root,
            "确认目录选对了；.xls 老格式请先另存为 .xlsx。",
        )
    return files


def unique_path(directory, filename):
    directory = Path(directory)
    target = directory / filename
    if not target.exists():
        return target
    stem, suffix = target.stem, target.suffix
    n = 1
    while (directory / ("%s_%d%s" % (stem, n, suffix))).exists():
        n += 1
    return directory / ("%s_%d%s" % (stem, n, suffix))


def safe_name(text, default="sheet"):
    name = re.sub(r'[\\/:*?"<>|\[\]]', "_", str(text)).strip()
    return name or default


# ── 单元格判断与表格读取 ──────────────────────────────────────────

def is_empty(value):
    return value is None or (isinstance(value, str) and value.strip() == "")


def trim_text(value):
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def to_number(value):
    """数值列判定：数字原样；纯数字字符串（可带千分位）转数字；其余不算。"""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        text = value.strip()
        if NUM_RE.fullmatch(text):
            try:
                return float(text.replace(",", ""))
            except ValueError:
                return None
    return None


def trim_trailing(rows):
    while rows and all(is_empty(v) for v in rows[-1]):
        rows.pop()
    width = 0
    for row in rows:
        for i, value in enumerate(row):
            if not is_empty(value):
                width = max(width, i + 1)
    return [(row + [None] * width)[:width] for row in rows]


def read_sheet_rows(ws):
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    return trim_trailing(rows)


def split_header(rows):
    """第一条非空行当表头，其余当数据。"""
    for idx, row in enumerate(rows):
        if any(not is_empty(v) for v in row):
            return row, rows[idx + 1:]
    return [], []


def header_names(header):
    names, used = [], {}
    for i, value in enumerate(header):
        base = str(value).strip() if not is_empty(value) else "列%d" % (i + 1)
        count = used.get(base, 0) + 1
        used[base] = count
        names.append(base if count == 1 else "%s#%d" % (base, count))
    return names


# ── 各动作 ────────────────────────────────────────────────────────

def action_merge(payload):
    files = collect_xlsx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    out.mkdir(parents=True, exist_ok=True)

    sheets = []  # (file, sheet_name, names, data_rows)
    union, seen = [], set()
    for src in files:
        wb = load_workbook(str(src), data_only=True)
        for ws in wb.worksheets:
            rows = read_sheet_rows(ws)
            header, data = split_header(rows)
            if not header:
                continue
            names = header_names(header)
            data = [r for r in data if any(not is_empty(v) for v in r)]
            sheets.append((src, ws.title, names, data))
            for name in names:
                if name not in seen:
                    seen.add(name)
                    union.append(name)

    header_row = [SOURCE_FILE, SOURCE_SHEET] + union
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "合并"
    ws_out.append(header_row)
    for cell in ws_out[1]:
        cell.font = Font(bold=True)

    total_rows = 0
    file_results = []
    for src, sheet_name, names, data in sheets:
        for row in data:
            mapping = dict(zip(names, row))
            ws_out.append([src.name, sheet_name] + [mapping.get(name) for name in union])
            total_rows += 1
        file_results.append({
            "input": str(src),
            "input_name": src.name,
            "sheet": sheet_name,
            "rows": len(data),
        })

    ws_out.freeze_panes = "A2"
    if ws_out.max_row >= 2:
        ws_out.auto_filter.ref = ws_out.dimensions
    target = unique_path(out, "merged.xlsx")
    wb_out.save(str(target))

    return {
        "action": "merge",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": file_results,
        "output": str(target),
        "output_name": target.name,
        "rows": total_rows,
        "sheets": len(sheets),
        "summary": "合并 %d 个文件 / %d 个工作表，写入 %d 行数据" % (len(files), len(sheets), total_rows),
        "artifacts": [str(target)],
    }


def normalize_rows(rows):
    """trim 文本 + 统计被 trim 的单元格数。"""
    trimmed, out = 0, []
    for row in rows:
        new_row = []
        for value in row:
            if isinstance(value, str):
                stripped = value.strip()
                if stripped != value:
                    trimmed += 1
                new_row.append(stripped if stripped else None)
            else:
                new_row.append(value)
        out.append(new_row)
    return out, trimmed


def action_clean(payload):
    files = collect_xlsx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    out.mkdir(parents=True, exist_ok=True)
    results = []
    for src in files:
        wb = load_workbook(str(src), data_only=True)
        wb_out = Workbook()
        wb_out.remove(wb_out.active)
        total_in = total_empty = total_dup = total_trim = total_out = 0
        for ws in wb.worksheets:
            rows = read_sheet_rows(ws)
            header, data = split_header(rows)
            header, trim_h = normalize_rows([header] if header else [])
            header = header[0] if header else []
            data, trim_d = normalize_rows(data)
            total_trim += trim_h + trim_d

            seen, kept = set(), []
            empty = dup = 0
            for row in data:
                if all(is_empty(v) for v in row):
                    empty += 1
                    continue
                key = tuple(row)
                if key in seen:
                    dup += 1
                    continue
                seen.add(key)
                kept.append(row)

            total_in += len(data)
            total_empty += empty
            total_dup += dup
            total_out += len(kept)

            new_ws = wb_out.create_sheet(title=safe_name(ws.title)[:31])
            if header:
                new_ws.append(header)
                for cell in new_ws[1]:
                    cell.font = Font(bold=True)
            for row in kept:
                new_ws.append(row)
            new_ws.freeze_panes = "A2"

        target = unique_path(out, src.stem + ".xlsx")
        wb_out.save(str(target))
        results.append({
            "input": str(src),
            "input_name": src.name,
            "output": str(target),
            "output_name": target.name,
            "rows_in": total_in,
            "rows_out": total_out,
            "empty_removed": total_empty,
            "duplicates_removed": total_dup,
            "cells_trimmed": total_trim,
        })
    return {
        "action": "clean",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": results,
        "summary": "清洗 %d 个文件：删空行 %d、去重 %d、trim %d 个单元格" % (
            len(files),
            sum(r["empty_removed"] for r in results),
            sum(r["duplicates_removed"] for r in results),
            sum(r["cells_trimmed"] for r in results),
        ),
        "artifacts": [r["output"] for r in results],
    }


def action_summary(payload):
    files = collect_xlsx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    out.mkdir(parents=True, exist_ok=True)
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "汇总"
    ws_out.append(["文件", "工作表", "列名", "非空计数", "数值个数", "求和"])
    for cell in ws_out[1]:
        cell.font = Font(bold=True)
    report_rows = 0
    file_results = []

    for src in files:
        wb = load_workbook(str(src), data_only=True)
        sheet_count = 0
        for ws in wb.worksheets:
            rows = read_sheet_rows(ws)
            header, data = split_header(rows)
            if not header:
                continue
            names = header_names(header)
            sheet_count += 1
            width = len(header)
            for col in range(width):
                values = [row[col] for row in data if col < len(row)]
                values = [v for v in values if not is_empty(v)]
                numbers = [n for n in (to_number(v) for v in values) if n is not None]
                ws_out.append([
                    src.name,
                    ws.title,
                    names[col],
                    len(values),
                    len(numbers),
                    sum(numbers) if numbers else None,
                ])
                report_rows += 1
        file_results.append({
            "input": str(src),
            "input_name": src.name,
            "sheets": sheet_count,
        })

    ws_out.freeze_panes = "A2"
    if ws_out.max_row >= 2:
        ws_out.auto_filter.ref = ws_out.dimensions
    for letter, width in (("A", 26), ("B", 22), ("C", 24), ("D", 12), ("E", 12), ("F", 14)):
        ws_out.column_dimensions[letter].width = width
    target = unique_path(out, "summary.xlsx")
    wb_out.save(str(target))

    return {
        "action": "summary",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": file_results,
        "output": str(target),
        "output_name": target.name,
        "rows": report_rows,
        "summary": "汇总 %d 个文件，输出 %d 行列级统计" % (len(files), report_rows),
        "artifacts": [str(target)],
    }


def action_csv(payload):
    files = collect_xlsx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    out.mkdir(parents=True, exist_ok=True)
    results = []
    for src in files:
        wb = load_workbook(str(src), data_only=True)
        for ws in wb.worksheets:
            rows = read_sheet_rows(ws)
            filename = "%s__%s.csv" % (safe_name(src.stem), safe_name(ws.title))
            target = unique_path(out, filename)
            with target.open("w", encoding="utf-8-sig", newline="") as fh:
                writer = csv.writer(fh)
                for row in rows:
                    writer.writerow(["" if v is None else v for v in row])
            results.append({
                "input": str(src),
                "input_name": src.name,
                "sheet": ws.title,
                "output": str(target),
                "output_name": target.name,
                "rows": len(rows),
            })
    return {
        "action": "csv",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": results,
        "summary": "导出 %d 个 CSV（UTF-8 带 BOM，Excel 可直接打开）" % len(results),
        "artifacts": [r["output"] for r in results],
    }


# ── demo ─────────────────────────────────────────────────────────

def _demo_book(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    wb.save(str(path))


def action_demo(payload):
    root = Path(tempfile.mkdtemp(prefix="dsh-xlsx-batch-demo-"))
    inputs = root / "样例"
    inputs.mkdir(parents=True, exist_ok=True)
    _demo_book(inputs / "销售A.xlsx", [
        ("销售", [
            ["姓名", "部门", "金额"],
            ["张三", " 销售部 ", "100"],
            ["李四", "销售部", "200"],
            [None, None, None],
            ["张三", " 销售部 ", "100"],
        ]),
        ("库存", [["品名", "数量"], ["A", 5], ["B", 7]]),
    ])
    _demo_book(inputs / "销售B.xlsx", [
        ("销售", [["部门", "姓名", "金额"], ["市场部", "王五", "300"]]),
        ("备注", [["说明"], ["样例数据"]]),
    ])
    out = root / "输出"
    steps, artifacts = [], []

    def run(name, fn):
        result = fn({"input": str(inputs), "out": str(out / name)})
        steps.append({"step": name, "action": result.get("action"),
                      "summary": result.get("summary", ""), "out": result.get("out")})
        artifacts.extend(result.get("artifacts", []))

    run("合并", action_merge)
    run("清洗", action_clean)
    run("汇总", action_summary)
    run("CSV", action_csv)

    return {
        "action": "demo",
        "demoRoot": str(root),
        "steps": steps,
        "summary": "演示完成：4 个动作全部跑通，产物 %d 个" % len(artifacts),
        "artifacts": artifacts,
    }


ACTIONS = {
    "merge": action_merge,
    "clean": action_clean,
    "summary": action_summary,
    "csv": action_csv,
    "demo": action_demo,
}


def main(argv):
    if len(argv) < 3:
        emit({"ok": False, "code": 2,
              "message": "内部调用参数不足：需要 <action> 和 JSON payload",
              "hint": "请通过 bin/cli.mjs 调用本脚本。"})
        return 2
    action, raw = argv[1], argv[2]
    try:
        payload = json.loads(raw) if raw else {}
    except json.JSONDecodeError as exc:
        emit({"ok": False, "code": 2, "message": "参数 JSON 解析失败：%s" % exc, "hint": ""})
        return 2
    fn = ACTIONS.get(action)
    if not fn:
        emit({"ok": False, "code": 2, "message": "未知动作：%s" % action,
              "hint": "可用动作：%s" % " / ".join(k for k in ACTIONS if k != "demo")})
        return 2
    try:
        result = fn(payload)
        result.setdefault("ok", True)
        emit(result)
        return 0
    except ToolError as exc:
        emit({"ok": False, "code": exc.code, "message": exc.message, "hint": exc.hint})
        return exc.code
    except PermissionError as exc:
        emit({"ok": False, "code": 1,
              "message": "没有权限读写文件：%s" % exc,
              "hint": "检查目标文件是否被 Excel 占用、或输出目录是否有写权限。"})
        return 1
    except Exception as exc:  # noqa: BLE001 —— 兜底，保证永远输出 JSON
        traceback.print_exc(file=sys.stderr)
        emit({"ok": False, "code": 1, "message": "处理失败：%s" % exc,
              "hint": "文件可能已损坏、被加密或格式不是标准 xlsx；可先用 Excel 另存一次再试。"})
        return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
